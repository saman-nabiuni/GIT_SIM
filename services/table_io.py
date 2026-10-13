"""Read/write tables: CSV (with delimiter and encoding detection) and XLSX.
The result is always (columns, rows); rows is a list of dicts and all values are str."""
import csv
import io
import os
import tempfile
from datetime import date, datetime
from pathlib import Path

from . import config
from .errors import ServiceError


def _clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    text = str(value).strip()
    if len(text) > config.MAX_CELL_LEN:
        raise ServiceError(f"طول یک سلول بیشتر از {config.MAX_CELL_LEN} کاراکتر است.")
    return text


def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "cp1256"):  # cp1256 = Excel's "CSV" save on Persian Windows
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ServiceError("انکودینگ فایل قابل خواندن نیست؛ آن را با UTF-8 ذخیره کن.")


def _sniff_delimiter(text: str) -> str:
    first = next((ln for ln in text.splitlines() if ln.strip()), "")
    counts = {d: first.count(d) for d in (",", "\t", ";")}
    best = max(counts, key=lambda d: (counts[d], d == ","))
    return best if counts[best] else ","


def _build(matrix) -> tuple:
    """matrix: list of rows (each a list of values). First row = header."""
    matrix = [r for r in matrix if any(str(c).strip() for c in r if c is not None)]
    if not matrix:
        raise ServiceError("فایل خالی است.")
    header = [_clean(c) for c in matrix[0]]
    body = [[_clean(c) for c in r] for r in matrix[1:]]

    # trailing columns with no name and no data are dropped
    while header and not header[-1] and all(not (r[len(header) - 1] if len(r) >= len(header) else "") for r in body):
        header.pop()
    if not header:
        raise ServiceError("سطر اول فایل باید نام ستون‌ها باشد.")
    if any(not h for h in header):
        raise ServiceError("یکی از ستون‌ها نام ندارد.")
    if len(set(h.lower() for h in header)) != len(header):
        raise ServiceError("نام ستون‌ها تکراری است.")
    if len(header) > config.MAX_TABLE_COLS:
        raise ServiceError(f"تعداد ستون‌ها از {config.MAX_TABLE_COLS} بیشتر است.")
    if len(body) > config.MAX_TABLE_ROWS:
        raise ServiceError(f"تعداد سطرها از {config.MAX_TABLE_ROWS} بیشتر است.")

    rows = []
    for i, r in enumerate(body, start=2):
        if any(r[len(header):]):
            raise ServiceError(f"سطر {i} بیشتر از تعداد ستون‌ها مقدار دارد.")
        r = (r + [""] * len(header))[: len(header)]
        rows.append(dict(zip(header, r)))
    return header, rows


def read_csv_bytes(raw: bytes) -> tuple:
    text = _decode(raw)
    delimiter = _sniff_delimiter(text)
    try:
        matrix = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    except csv.Error as exc:
        raise ServiceError(f"CSV معتبر نیست: {exc}")
    return _build(matrix)


def read_xlsx_bytes(raw: bytes) -> tuple:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ServiceError("برای خواندن xlsx باید کتابخانه‌ی openpyxl نصب باشد: pip install openpyxl", 500)
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        matrix = []
        for row in ws.iter_rows(values_only=True):
            matrix.append(list(row))
            if len(matrix) > config.MAX_TABLE_ROWS + 2:
                raise ServiceError(f"تعداد سطرها از {config.MAX_TABLE_ROWS} بیشتر است.")
        wb.close()
    except ServiceError:
        raise
    except Exception:
        raise ServiceError("فایل xlsx خراب است یا قابل خواندن نیست.")
    return _build(matrix)


def read_upload(file_storage) -> tuple:
    """Converts an uploaded file (CSV or XLSX) to (columns, rows)."""
    name = (file_storage.filename or "").lower()
    raw = file_storage.read()
    if not raw:
        raise ServiceError("فایل خالی است.")
    if name.endswith(".xlsx"):
        return read_xlsx_bytes(raw)
    if name.endswith(".csv") or name.endswith(".txt"):
        return read_csv_bytes(raw)
    if name.endswith(".xls"):
        raise ServiceError("فرمت xls قدیمی پشتیبانی نمی‌شود؛ فایل را به xlsx یا csv تبدیل کن.")
    raise ServiceError("فقط فایل csv یا xlsx قابل قبول است.")


def read_csv_path(path: Path) -> tuple:
    return read_csv_bytes(Path(path).read_bytes())


def write_csv(path: Path, columns, rows) -> None:
    """Atomic write with utf-8-sig so Persian Excel opens it correctly."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    buf = io.StringIO(newline="")
    writer = csv.writer(buf)
    writer.writerow(columns)
    for row in rows:
        writer.writerow([row.get(c, "") for c in columns])
    atomic_write_bytes(path, buf.getvalue().encode("utf-8-sig"))


def atomic_write_bytes(path: Path, data: bytes) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(dir=str(path.parent), prefix=path.name + ".", suffix=".tmp")
    try:
        with os.fdopen(fd, "wb") as fh:
            fh.write(data)
        os.replace(tmp, path)
    except Exception:
        try:
            os.unlink(tmp)
        except OSError:
            pass
        raise
